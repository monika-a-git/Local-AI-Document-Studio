"""
Multi-format export manager for Local AI Document Studio.
Exports generated content to PDF (ReportLab), DOCX (python-docx),
Markdown, TXT, CSV, and PNG.
"""

import io
import re
import csv
from datetime import datetime
from typing import Optional, Dict, Any, List, Union
from pathlib import Path
from PIL import Image
import pandas as pd

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.pdfgen import canvas

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.config import EXPORTS_DIR


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically calculate and render 'Page X of Y' in footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748B"))
        # Header rule
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, letter[1] - 40, letter[0] - 54, letter[1] - 40)
        self.drawString(54, letter[1] - 34, "Local AI Document Studio — Confidential & Private")

        # Footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 30, page_str)
        self.drawString(54, 30, "Generated locally with AI Studio")
        self.line(54, 42, letter[0] - 54, 42)
        self.restoreState()


def export_to_pdf(
    title: str,
    markdown_content: str,
    output_path: Optional[Union[str, Path]] = None,
    metadata: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Generate professional PDF document using ReportLab.
    Includes custom headers, styled headings, bullet points, tables, and page numbers.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=10,
    )

    meta_style = ParagraphStyle(
        "MetaStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=15,
    )

    h1_style = ParagraphStyle(
        "H1Style",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )

    h2_style = ParagraphStyle(
        "H2Style",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#2563EB"),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=8,
    )

    bullet_style = ParagraphStyle(
        "BulletStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1E293B"),
        leftIndent=15,
        spaceAfter=4,
    )

    story = []

    # Title & Metadata
    clean_title = title.replace("#", "").strip()
    story.append(Paragraph(clean_title, title_style))

    meta_text = f"Date: {datetime.now().strftime('%B %d, %Y')} | Generated locally with Ollama & ChromaDB"
    if metadata and "source" in metadata:
        meta_text += f" | Source: {metadata['source']}"
    story.append(Paragraph(meta_text, meta_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3B82F6"), spaceAfter=15))

    # Parse markdown content into ReportLab flowables
    lines = markdown_content.splitlines()
    table_buffer = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if table_buffer:
                story.append(_render_pdf_table(table_buffer, body_style))
                table_buffer = []
                story.append(Spacer(1, 10))
            continue

        # Check for Markdown table line
        if stripped.startswith("|") and stripped.endswith("|"):
            table_buffer.append(stripped)
            continue
        elif table_buffer:
            story.append(_render_pdf_table(table_buffer, body_style))
            table_buffer = []
            story.append(Spacer(1, 10))

        # Check for Headings
        if stripped.startswith("# "):
            story.append(Paragraph(stripped[2:], h1_style))
        elif stripped.startswith("## "):
            story.append(Paragraph(stripped[3:], h1_style))
        elif stripped.startswith("### "):
            story.append(Paragraph(stripped[4:], h2_style))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            clean_item = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", stripped[2:])
            story.append(Paragraph(f"• {clean_item}", bullet_style))
        elif re.match(r"^\d+\.\s+", stripped):
            clean_item = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", stripped)
            story.append(Paragraph(clean_item, bullet_style))
        elif stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            quote_style = ParagraphStyle(
                "Quote",
                parent=body_style,
                fontName="Helvetica-Oblique",
                textColor=colors.HexColor("#475569"),
                leftIndent=20,
            )
            story.append(Paragraph(quote_text, quote_style))
        else:
            clean_p = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", stripped)
            story.append(Paragraph(clean_p, body_style))

    if table_buffer:
        story.append(_render_pdf_table(table_buffer, body_style))

    # Build PDF
    doc.build(story, canvasmaker=NumberedCanvas)
    pdf_bytes = buffer.getvalue()
    buffer.close()

    if output_path:
        Path(output_path).write_bytes(pdf_bytes)

    return pdf_bytes


def _render_pdf_table(table_lines: List[str], body_style: ParagraphStyle) -> Table:
    """Helper to convert Markdown table lines to a styled ReportLab Table."""
    rows = []
    for line in table_lines:
        if re.match(r"^\|[\s\-:\s|]+\|$", line):
            continue  # Separator line like |---|---|
        cells = [c.strip() for c in line.strip("|").split("|")]
        # Convert cells to Paragraphs for word wrapping
        flowable_cells = [Paragraph(re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", c), body_style) for c in cells]
        rows.append(flowable_cells)

    if not rows:
        return Table([["No data"]])

    t = Table(rows, colWidths=None)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
    ]))
    return t


def export_to_docx(
    title: str,
    markdown_content: str,
    output_path: Optional[Union[str, Path]] = None,
    metadata: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Generate professional Microsoft Word (.docx) document using python-docx.
    """
    doc = docx.Document()

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title.replace("#", "").strip())
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

    # Subtitle / Metadata
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        f"Generated locally on {datetime.now().strftime('%B %d, %Y')} | Local AI Document Studio"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Content Lines
    lines = markdown_content.splitlines()
    table_lines = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            if table_lines:
                _render_docx_table(doc, table_lines)
                table_lines = []
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines.append(stripped)
            continue
        elif table_lines:
            _render_docx_table(doc, table_lines)
            table_lines = []

        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            h.runs[0].font.color.rgb = RGBColor(37, 99, 235)
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
            h.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs(p, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            clean_text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _add_formatted_runs(p, clean_text)
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            run = p.add_run(stripped.lstrip("> ").strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        else:
            p = doc.add_paragraph()
            _add_formatted_runs(p, stripped)

    if table_lines:
        _render_docx_table(doc, table_lines)

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()

    if output_path:
        Path(output_path).write_bytes(docx_bytes)

    return docx_bytes


def _add_formatted_runs(paragraph, text: str):
    """Parse inline bold **text** and add runs."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
        else:
            paragraph.add_run(part)


def _render_docx_table(doc, table_lines: List[str]):
    """Helper to convert Markdown table lines to docx Table."""
    parsed_rows = []
    for line in table_lines:
        if re.match(r"^\|[\s\-:\s|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_cols = max(len(r) for r in parsed_rows)
    table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
    table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else "Table Grid"

    for r_idx, row in enumerate(parsed_rows):
        for c_idx, cell_value in enumerate(row):
            if c_idx < num_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = cell_value
                if r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True


def export_to_markdown(
    title: str,
    content: str,
    output_path: Optional[Union[str, Path]] = None,
    metadata: Optional[Dict[str, Any]] = None,
) -> str:
    """Format and export clean markdown with YAML-style frontmatter."""
    date_str = datetime.now().isoformat()
    meta_frontmatter = (
        f"---\n"
        f"title: \"{title}\"\n"
        f"generated_at: \"{date_str}\"\n"
        f"tool: \"Local AI Document Studio\"\n"
        f"---\n\n"
    )
    full_md = meta_frontmatter + content.strip() + "\n"

    if output_path:
        Path(output_path).write_text(full_md, encoding="utf-8")

    return full_md


def export_to_txt(
    title: str,
    content: str,
    output_path: Optional[Union[str, Path]] = None,
) -> str:
    """Format and export plain text."""
    clean_text = re.sub(r"[\#\*\`]", "", content).strip()
    header = f"{title.upper()}\n{'=' * len(title)}\nGenerated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n"
    full_txt = header + clean_text + "\n"

    if output_path:
        Path(output_path).write_text(full_txt, encoding="utf-8")

    return full_txt


def export_to_csv(
    markdown_or_table_data: Union[str, List[Dict[str, Any]]],
    output_path: Optional[Union[str, Path]] = None,
) -> str:
    """
    Extract table rows from markdown or list of dicts and export to CSV string.
    """
    if isinstance(markdown_or_table_data, str):
        # Extract markdown table lines
        lines = markdown_or_table_data.splitlines()
        table_rows = []
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("|") and stripped.endswith("|"):
                if re.match(r"^\|[\s\-:\s|]+\|$", stripped):
                    continue
                cells = [c.strip() for c in stripped.strip("|").split("|")]
                table_rows.append(cells)

        output = io.StringIO()
        writer = csv.writer(output)
        for row in table_rows:
            writer.writerow(row)
        csv_str = output.getvalue()
    elif isinstance(markdown_or_table_data, list):
        df = pd.DataFrame(markdown_or_table_data)
        csv_str = df.to_csv(index=False)
    else:
        csv_str = ""

    if output_path:
        Path(output_path).write_text(csv_str, encoding="utf-8")

    return csv_str


def export_image_to_png(image: Image.Image, output_path: Optional[Union[str, Path]] = None) -> bytes:
    """Convert PIL Image to PNG bytes."""
    buf = io.BytesIO()
    image.save(buf, format="PNG", dpi=(150, 150))
    png_bytes = buf.getvalue()
    buf.close()

    if output_path:
        Path(output_path).write_bytes(png_bytes)

    return png_bytes

